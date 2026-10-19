"""Generate downloads from exactly the rows and fields visible in the viewer."""
import io
import math
from datetime import datetime
from xml.sax.saxutils import escape
from pathlib import Path
from documents import parse_date

NUMERIC = {"Transaction amount", "Fare", "Net sales", "Commission", "Discount",
           "Tax on commission", "Balance payable", "Page"}
ALLOWED = NUMERIC | {"Agency", "Agency code", "Type", "Document number", "Date",
                     "Taxes (codes)", "Fees (codes)", "Penalties (codes)",
                     "Payment code", "Related document", "Notes"}


def validate(payload):
    if not isinstance(payload, dict):
        raise ValueError("Invalid export selection.")
    fields, rows = payload.get("fields"), payload.get("rows")
    available = payload.get("available_fields", ALLOWED)
    if (not isinstance(fields, list) or not fields or len(fields) > len(available)
            or any(not isinstance(f, str) or f not in available for f in fields)
            or len(set(fields)) != len(fields)):
        raise ValueError("Select at least one valid field to export.")
    if not isinstance(rows, list) or not rows:
        raise ValueError("There are no matching records to export.")
    if len(rows) > 1_048_568:
        raise ValueError("This selection exceeds Excel's row limit. Filter it before exporting.")
    for row in rows:
        if not isinstance(row, dict):
            raise ValueError("Invalid record in export selection.")
        for field in fields:
            value = row.get(field)
            if value is not None and not isinstance(value, (str, int, float)):
                raise ValueError("Invalid value in export selection.")
            if isinstance(value, str) and len(value) > 32767:
                raise ValueError("A selected value is too long to export.")
            if isinstance(value, float) and not math.isfinite(value):
                raise ValueError("Invalid number in export selection.")
    return fields, rows


def context(payload, count):
    return [str(payload.get("filename", "Uploaded PDF"))[:250],
            str(payload.get("period", ""))[:250],
            str(payload.get("filters", "All records"))[:1000],
            f"{count} selected records." + (" Amounts in INR. Blank amounts are unavailable." if payload.get("currency")=="INR" else "")]


def excel(payload, fields, rows):
    # Runtime export uses a portable Python library; no Office installation needed.
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Selected records"
    sheet.append([payload.get("title", "Selected PDF information")])
    for line in context(payload, len(rows)):
        sheet.append([line])
    sheet.append([])
    sheet.append(fields)
    header = sheet.max_row
    for row in rows:
        sheet.append([row.get(field) for field in fields])
    # Explicit text typing preserves identifiers/leading zeros and prevents formulas
    # from document text being executed by Excel.
    for row in sheet:
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = "s"
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for col, field in enumerate(fields, 1):
        label = sheet.cell(header, col)
        label.font = Font(bold=True, color="FFFFFF")
        label.fill = PatternFill("solid", fgColor="245EDB")
        sheet.column_dimensions[get_column_letter(col)].width = (
            70 if field=="Text" else 44 if field in {"Agency", "Notes"} else 28 if "(codes)" in field else 23)
        for r in range(header + 1, sheet.max_row + 1):
            cell = sheet.cell(r, col)
            kind = payload.get("field_types", {}).get(field, "number" if field in NUMERIC else "date" if field=="Date" else "text")
            if kind=="number" and isinstance(cell.value, (int, float)):
                cell.number_format = '#,##0.##;[Red]-#,##0.##;0'
            elif kind=="date" and cell.value:
                value = parse_date(cell.value)
                if value:
                    cell.value = value
                    cell.number_format = "dd-mmm-yyyy"
    # Metadata lives above the table without stretching the selected columns.
    if len(fields) > 1:
        for r in range(1, 6):
            sheet.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(fields))
    sheet["A1"].font = Font(size=16, bold=True, color="17243A")
    for r in range(1, 6):
        sheet.row_dimensions[r].height = 32 if r == 4 else 25
    sheet.freeze_panes = f"A{header + 1}"
    sheet.auto_filter.ref = f"A{header}:{get_column_letter(len(fields))}{sheet.max_row}"
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def pdf(payload, fields, rows):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_RIGHT
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, LongTable, TableStyle, PageBreak
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    font = "Helvetica"
    for font_path in (Path(r"C:\Windows\Fonts\arial.ttf"), Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf")):
        if font_path.exists():
            if "ViewerText" not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont("ViewerText",str(font_path)))
            font = "ViewerText"
            break
    output = io.BytesIO()
    page_size = landscape(A4)
    doc = SimpleDocTemplate(output, pagesize=page_size, rightMargin=28,
                            leftMargin=28, topMargin=28, bottomMargin=32,
                            title=payload.get("title","Selected PDF information"))
    styles = getSampleStyleSheet()
    normal = ParagraphStyle("Cell", fontName=font, fontSize=8,
                            leading=10, splitLongWords=True)
    numeric = ParagraphStyle("Number", parent=normal, alignment=TA_RIGHT)
    head = ParagraphStyle("Header", parent=normal, textColor=colors.white)
    def p(value, style=normal):
        return Paragraph(escape("" if value is None else str(value)).replace("\n", "<br/>"), style)
    story = [Paragraph(escape(payload.get("title","Selected PDF information")), styles["Title"])]
    for line in context(payload, len(rows)):
        story.extend([p(line), Spacer(1, 5)])
    # Split wide selections into readable column groups, keeping the same record
    # numbering in every group. Long rows can split across pages without clipping.
    groups = [fields[i:i + 5] for i in range(0, len(fields), 5)]
    for group_no, group in enumerate(groups, 1):
        if group_no > 1:
            story.append(PageBreak())
        story.extend([Spacer(1, 10), p(f"Fields {group_no} of {len(groups)}"), Spacer(1, 7)])
        table_rows = [[p("Row", head)] + [p(f, head) for f in group]]
        for index, row in enumerate(rows, 1):
            cells = [p(index)]
            for field in group:
                value = row.get(field)
                kind = payload.get("field_types", {}).get(field,"number" if field in NUMERIC else "text")
                is_number = kind=="number" and isinstance(value, (float, int))
                cells.append(p(f"{value:,.2f}".removesuffix(".00") if is_number else value,
                               numeric if is_number else normal))
            table_rows.append(cells)
        table = LongTable(table_rows, colWidths=[32] + [(doc.width - 32) / len(group)] * len(group),
                          repeatRows=1, splitByRow=1, splitInRow=1, hAlign="LEFT")
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#245EDB")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F1F5FB")]),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 6),
            ("RIGHTPADDING", (0, 0), (-1, -1), 6),
            ("TOPPADDING", (0, 0), (-1, -1), 6),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ("LINEBELOW", (0, 0), (-1, 0), .5, colors.HexColor("#245EDB")),
        ]))
        story.append(table)
    def footer(canvas, doc):
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.HexColor("#627089"))
        canvas.drawRightString(page_size[0] - 28, 17, f"Page {doc.page}")
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return output.getvalue()


def export_selection(payload, fmt):
    fields, rows = validate(payload)
    if fmt == "xlsx":
        return excel(payload, fields, rows), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    if fmt == "pdf":
        return pdf(payload, fields, rows), "application/pdf"
    raise ValueError("Choose PDF or Excel export.")
